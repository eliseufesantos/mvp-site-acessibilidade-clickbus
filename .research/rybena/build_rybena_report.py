from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / ".research" / "rybena" / "report-source.md"
OUTPUT = ROOT / "docs" / "Pesquisa_Rybena_FIAP.docx"

INK = "161B22"
MUTED = "5B6573"
ACCENT = "16697A"
ACCENT_DARK = "0B3C49"
ACCENT_LIGHT = "E9F4F6"
BLUE_LIGHT = "EDF4FA"
GRAY_LIGHT = "F3F5F7"
GRAY_BORDER = "C7CDD4"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRAY_BORDER, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, sep, text, end])


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([color, underline])
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, *, color: str | None = None, size: float | None = None) -> None:
    cursor = 0
    for match in INLINE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            if size:
                run.font.size = Pt(size)
        token = match.group(0)
        if token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            if size:
                run.font.size = Pt(size)
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.2 if size is None else size)
            run.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)


def apply_run_defaults(paragraph) -> None:
    for run in paragraph.runs:
        if not run.font.name:
            run.font.name = "Aptos"
        if not run.font.size:
            run.font.size = Pt(10.3)


def no_split_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = {
        2: [2.05, 4.85],
        3: [1.55, 2.75, 2.60],
        4: [1.35, 2.25, 1.35, 1.95],
    }.get(cols, [6.9 / cols] * cols)
    for row_index, (row, values) in enumerate(zip(table.rows, rows)):
        no_split_row(row)
        if row_index == 0:
            repeat_table_header(row)
        for col_index, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, ACCENT_DARK)
            elif row_index % 2 == 0:
                set_cell_shading(cell, BLUE_LIGHT)
            else:
                set_cell_shading(cell, WHITE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value, color=WHITE if row_index == 0 else INK, size=8.6)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_body_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Body Text")
    segments = text.split("\n")
    for index, segment in enumerate(segments):
        if index:
            p.add_run().add_break(WD_BREAK.LINE)
        add_inline(p, segment)
    apply_run_defaults(p)


def add_bullet(doc: Document, text: str, numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    p = doc.add_paragraph(style=style)
    add_inline(p, text)
    apply_run_defaults(p)


def add_numbered(doc: Document, number: str, text: str) -> None:
    p = doc.add_paragraph(style="Body Text")
    p.paragraph_format.left_indent = Inches(0.34)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    prefix = p.add_run(f"{number}.  ")
    prefix.bold = True
    prefix.font.name = "Aptos"
    prefix.font.size = Pt(10.1)
    add_inline(p, text)
    apply_run_defaults(p)


def add_heading(doc: Document, text: str, level: int, section_number: int | None = None) -> None:
    if level == 2 and section_number is not None:
        text = f"{section_number:02d}  {text}"
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    add_inline(p, text)


def add_cover(doc: Document, title: str, subtitle: str, metadata: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("FIAP  ·  CLICKBUS")
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(ACCENT)

    bar = doc.add_table(rows=1, cols=1)
    bar.autofit = False
    repeat_table_header(bar.rows[0])
    bar.columns[0].width = Inches(1.15)
    cell = bar.cell(0, 0)
    set_cell_shading(cell, ACCENT)
    set_cell_border(cell, ACCENT, "0")
    cell.height = Inches(0.08)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(120)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Aptos Display"
    r.font.size = Pt(30)
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(32)
    r = p.add_run(subtitle)
    r.font.name = "Aptos"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(MUTED)

    callout = doc.add_table(rows=1, cols=1)
    callout.autofit = False
    repeat_table_header(callout.rows[0])
    callout.columns[0].width = Inches(6.9)
    cell = callout.cell(0, 0)
    set_cell_shading(cell, ACCENT_LIGHT)
    set_cell_border(cell, ACCENT_LIGHT, "0")
    set_cell_margins(cell, top=180, start=220, bottom=180, end=220)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, "Benchmark funcional, análise crítica e especificação independente para um MVP acadêmico de comunicação acessível em viagens.", color=ACCENT_DARK, size=11)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(54)
    for index, line in enumerate(metadata):
        if index:
            p.add_run().add_break(WD_BREAK.LINE)
        add_inline(p, re.sub(r"\s{2,}$", "", line), color=MUTED, size=10)


def add_contents(doc: Document, headings: list[str]) -> None:
    doc.add_page_break()
    add_heading(doc, "Sumário", 1)
    p = doc.add_paragraph(style="Lead")
    add_inline(p, "Estrutura do relatório para leitura rápida. Os títulos são clicáveis apenas nas referências externas; a numeração corresponde às seções do documento.")
    cells = []
    numbered = [f"{i:02d}  {heading}" for i, heading in enumerate(headings, 1)]
    midpoint = (len(numbered) + 1) // 2
    for i in range(midpoint):
        left = numbered[i]
        right = numbered[i + midpoint] if i + midpoint < len(numbered) else ""
        cells.append([left, right])
    table = doc.add_table(rows=len(cells), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        no_split_row(row)
        for col_index, cell in enumerate(row.cells):
            cell.width = Inches(3.35)
            set_cell_margins(cell, top=115, start=140, bottom=115, end=140)
            set_cell_shading(cell, GRAY_LIGHT if row_index % 2 else WHITE)
            set_cell_border(cell, WHITE, "0")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, cells[row_index][col_index], color=INK, size=9.2)
    doc.add_page_break()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.1

    body = doc.styles["Body Text"]
    body.font.name = "Aptos"
    body.font.size = Pt(10.3)
    body.font.color.rgb = RGBColor.from_string(INK)
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.1
    body.paragraph_format.widow_control = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.1)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.left_indent = Inches(0.25)
        style.paragraph_format.first_line_indent = Inches(-0.16)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.widow_control = True

    lead = doc.styles.add_style("Lead", WD_STYLE_TYPE.PARAGRAPH)
    lead.font.name = "Aptos"
    lead.font.size = Pt(11.2)
    lead.font.color.rgb = RGBColor.from_string(MUTED)
    lead.paragraph_format.space_after = Pt(14)
    lead.paragraph_format.line_spacing = 1.12

    heading_specs = {
        1: (22, 20, 12),
        2: (16, 18, 8),
        3: (12.5, 12, 5),
    }
    for level, (size, before, after) in heading_specs.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("FIAP · ClickBus   |   ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_field(footer)


def parse_markdown(doc: Document, lines: list[str], body_start: int) -> None:
    i = body_start
    section_number = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 3)
            i += 1
            continue
        if line.startswith("## "):
            section_number += 1
            add_heading(doc, line[3:].strip(), 2, section_number)
            i += 1
            continue
        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), 1)
            i += 1
            continue
        if line.startswith("|"):
            raw_rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw_rows.append([part.strip() for part in lines[i].strip().strip("|").split("|")])
                i += 1
            if len(raw_rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in raw_rows[1]):
                raw_rows.pop(1)
            add_table(doc, raw_rows)
            continue
        if re.match(r"^\d+\.\s+", line):
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                match = re.match(r"^(\d+)\.\s+(.+)$", lines[i].strip())
                add_numbered(doc, match.group(1), match.group(2))
                i += 1
            continue
        if line.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                add_bullet(doc, lines[i].strip()[2:])
                i += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt or nxt == "<!-- pagebreak -->" or nxt.startswith(("# ", "## ", "### ", "|", "- ")) or re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        combined = ""
        for part in paragraph_lines:
            hard_break = part.endswith("  ")
            combined += part.rstrip() + ("\n" if hard_break else " ")
        add_body_paragraph(doc, combined.strip())


def main() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = lines[0][2:].strip()
    subtitle = lines[2][3:].strip()
    first_break = lines.index("<!-- pagebreak -->")
    metadata = [line for line in lines[4:first_break] if line.strip()]
    body_lines = lines[first_break + 1:]
    headings = [line[3:].strip() for line in body_lines if line.startswith("## ")]

    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = title
    props.subject = "Benchmark da Rybená e blueprint de MVP acadêmico para a FIAP"
    props.author = "Equipe ClickBus FIAP"
    props.keywords = "Rybená, Libras, acessibilidade, FIAP, ClickBus, benchmark, MVP"
    props.comments = "Elaborado a partir de fontes públicas consultadas em setembro de 2026."

    add_cover(doc, title, subtitle, metadata)
    add_contents(doc, headings)
    parse_markdown(doc, lines, first_break + 1)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
